"""Tests for reports service and Excel generation."""

import io
import pytest
import openpyxl

from gatebot.db.crud import (
    add_or_update_protected_group,
    add_or_update_required_channel,
    get_daily_join_stats,
    log_join_event,
)
from gatebot.services.reports import generate_excel_report


@pytest.mark.asyncio
async def test_get_daily_join_stats(db_session):
    """Test daily join request analytics aggregation."""
    g1 = await add_or_update_protected_group(db_session, chat_id=-100111, title="Test Group 1")
    g2 = await add_or_update_protected_group(db_session, chat_id=-100222, title="Test Group 2")

    # Add join events
    await log_join_event(db_session, group_id=g1.id, user_id=1, user_name="User1", status="approved")
    await log_join_event(db_session, group_id=g1.id, user_id=2, user_name="User2", status="declined", missing_channels="Channel 1")
    await log_join_event(db_session, group_id=g2.id, user_id=3, user_name="User3", status="approved")

    daily = await get_daily_join_stats(db_session, days=7)
    assert len(daily) >= 2
    # Verify group titles mapped
    titles = [d["group_title"] for d in daily]
    assert "Test Group 1" in titles
    assert "Test Group 2" in titles


@pytest.mark.asyncio
async def test_generate_excel_report(db_session):
    """Test generating a valid Excel workbook with openpyxl."""
    g = await add_or_update_protected_group(db_session, chat_id=-100111, title="VIP Group")
    await log_join_event(db_session, group_id=g.id, user_id=10, user_name="Ali", status="approved")
    await log_join_event(db_session, group_id=g.id, user_id=20, user_name="Vali", status="declined", missing_channels="News || Sport")

    excel_bytes = await generate_excel_report(db_session)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000

    # Load and inspect sheets
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Guruhlar Reytingi" in sheet_names
    assert "Kunlik Dinamika" in sheet_names
    assert "Yetishmagan Kanallar" in sheet_names

    ws1 = wb["Guruhlar Reytingi"]
    assert "AQL ZIYOSI" in ws1["A1"].value
    assert "VIP Group" in [ws1.cell(row=r, column=2).value for r in range(9, 15)]
