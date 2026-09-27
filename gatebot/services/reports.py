"""Excel report generator for join request statistics using openpyxl."""

import io
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from gatebot.db.crud import get_daily_join_stats, get_join_stats

logger = logging.getLogger(__name__)


def _style_header(cell, text: str) -> None:
    cell.value = text
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_border(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Side(border_style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border


def _auto_column_width(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


async def generate_excel_report(session: AsyncSession) -> bytes:
    """Generate multi-sheet Excel report workbook with formatting."""
    stats = await get_join_stats(session)
    daily = await get_daily_join_stats(session, days=60)

    wb = openpyxl.Workbook()

    # --- Sheet 1: Umumiy va Reyting ---
    ws1 = wb.active
    ws1.title = "Guruhlar Reytingi"

    # Title Banner
    ws1.merge_cells("A1:F1")
    title_cell = ws1["A1"]
    title_cell.value = "AQL ZIYOSI — GURUHLAR BO'YICHA STATISTIKA VA REYTING"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 30

    # Summary box
    ws1["A3"] = "Jami tasdiqlangan (All-time):"
    ws1["B3"] = stats["total_approved"]
    ws1["A4"] = "Jami rad etilgan (All-time):"
    ws1["B4"] = stats["total_declined"]
    ws1["A5"] = "Oxirgi 7 kun tasdiqlangan:"
    ws1["B5"] = stats["approved_7d"]
    ws1["A6"] = "Oxirgi 7 kun rad etilgan:"
    ws1["B6"] = stats["declined_7d"]

    for r in range(3, 7):
        ws1[f"A{r}"].font = Font(bold=True)
        ws1[f"B{r}"].alignment = Alignment(horizontal="right")

    # Table Header for Groups Ranking
    headers = ["O'rin", "Guruh nomi", "Holati", "Tasdiqlangan", "Rad etilgan", "Jami so'rovlar", "Qabul %"]
    row_idx = 8
    ws1.row_dimensions[row_idx].height = 25
    for col_idx, h in enumerate(headers, start=1):
        cell = ws1.cell(row=row_idx, column=col_idx)
        _style_header(cell, h)

    # Table Rows
    top_groups = stats.get("top_groups", [])
    start_row = row_idx + 1
    for rank, g in enumerate(top_groups, start=1):
        r = start_row + rank - 1
        total = g["total"]
        pct = f"{(g['approved'] / total * 100):.1f}%" if total > 0 else "0.0%"
        status_str = "Faol" if g["is_active"] else "Nofaol"

        ws1.cell(row=r, column=1, value=rank).alignment = Alignment(horizontal="center")
        ws1.cell(row=r, column=2, value=g["title"])
        ws1.cell(row=r, column=3, value=status_str).alignment = Alignment(horizontal="center")
        ws1.cell(row=r, column=4, value=g["approved"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=r, column=5, value=g["declined"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=r, column=6, value=total).alignment = Alignment(horizontal="right")
        ws1.cell(row=r, column=7, value=pct).alignment = Alignment(horizontal="center")

    if top_groups:
        _apply_border(ws1, min_row=row_idx, max_row=start_row + len(top_groups) - 1, min_col=1, max_col=7)

    _auto_column_width(ws1)

    # --- Sheet 2: Kunlik Dinamika ---
    ws2 = wb.create_sheet(title="Kunlik Dinamika")

    ws2.merge_cells("A1:E1")
    title2 = ws2["A1"]
    title2.value = "KUNLAR VA GURUHLAR KESIMIDA SO'ROVLAR HISOBOTI"
    title2.font = Font(name="Arial", size=13, bold=True, color="1F4E79")
    title2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    headers2 = ["Sana", "Guruh nomi", "Tasdiqlangan", "Rad etilgan", "Jami"]
    ws2.row_dimensions[3].height = 22
    for col_idx, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=3, column=col_idx)
        _style_header(cell, h)

    row2 = 4
    for item in daily:
        ws2.cell(row=row2, column=1, value=item["date"]).alignment = Alignment(horizontal="center")
        ws2.cell(row=row2, column=2, value=item["group_title"])
        ws2.cell(row=row2, column=3, value=item["approved"]).alignment = Alignment(horizontal="right")
        ws2.cell(row=row2, column=4, value=item["declined"]).alignment = Alignment(horizontal="right")
        ws2.cell(row=row2, column=5, value=item["total"]).alignment = Alignment(horizontal="right")
        row2 += 1

    if daily:
        _apply_border(ws2, min_row=3, max_row=row2 - 1, min_col=1, max_col=5)
    _auto_column_width(ws2)

    # --- Sheet 3: Yetishmagan Kanallar ---
    ws3 = wb.create_sheet(title="Yetishmagan Kanallar")
    headers3 = ["Kanal nomi", "Rad etilish soni"]
    ws3.row_dimensions[1].height = 22
    for col_idx, h in enumerate(headers3, start=1):
        cell = ws3.cell(row=1, column=col_idx)
        _style_header(cell, h)

    row3 = 2
    for ch_title, count in stats.get("top_missing", []):
        ws3.cell(row=row3, column=1, value=ch_title)
        ws3.cell(row=row3, column=2, value=count).alignment = Alignment(horizontal="right")
        row3 += 1

    if stats.get("top_missing"):
        _apply_border(ws3, min_row=1, max_row=row3 - 1, min_col=1, max_col=2)
    _auto_column_width(ws3)

    # Save to in-memory bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
